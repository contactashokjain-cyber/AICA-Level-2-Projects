from __future__ import annotations
import csv
from pathlib import Path

BOOK_FIELDS=['title','author','publisher','edition','isbn','language','page_count','category','publication_year','cover_path','acquisition_date','source','cost','remarks','accession_no','rack','shelf','compartment']
MEMBER_FIELDS=['member_id','sno','name','father_name','address','registration_no','mobile','registration_date','status','remarks']

def import_rows(path):
    path=Path(path)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True)
        ws=wb.active
        values=list(ws.iter_rows(values_only=True)); wb.close()
        if not values: return []
        headers=[str(x or '').strip().lower() for x in values[0]]
        return [dict(zip(headers,row)) for row in values[1:] if any(v is not None for v in row)]
    if path.suffix.lower()=='.xls':
        try: import xlrd
        except ImportError as e: raise RuntimeError('Legacy .xls import requires xlrd. Install it with: pip install xlrd') from e
        wb=xlrd.open_workbook(path); sh=wb.sheet_by_index(0)
        headers=[str(sh.cell_value(0,c)).strip().lower() for c in range(sh.ncols)]
        return [dict(zip(headers,[sh.cell_value(r,c) for c in range(sh.ncols)])) for r in range(1,sh.nrows)]
    if path.suffix.lower()=='.csv':
        with path.open(newline='',encoding='utf-8-sig') as f: return list(csv.DictReader(f))
    raise ValueError('Supported formats: .xlsx, .xls (optional dependency), and .csv')

def export_template(path, fields):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title='Import'
    ws.append(fields); ws.freeze_panes='A2'; wb.save(path)

def export_xlsx(path, rows, sheet_name='Report'):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title=sheet_name[:31]
    if rows:
        headers=list(rows[0].keys()); ws.append(headers)
        for row in rows: ws.append([row.get(h) for h in headers])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    wb.save(path)

def export_pdf(path,title,rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A4)); styles=getSampleStyleSheet(); story=[Paragraph(title,styles['Title']),Spacer(1,12)]
    if rows:
        headers=list(rows[0].keys()); data=[[str(x) for x in headers]]+[[str(row.get(h,'') if row.get(h) is not None else '')[:90] for h in headers] for row in rows]
        table=Table(data,repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),0.35,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP'),('FONTSIZE',(0,0),(-1,-1),7)]))
        story.append(table)
    else: story.append(Paragraph('No records found.',styles['Normal']))
    doc.build(story)
